"""Build the editable public resume from the same approved data as the website.

Requires python-docx. Export the resulting DOCX with export-resume.ps1 on Windows.
The private original is deliberately never read or copied by this generator.
"""
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "resume" / "isac-zarate-resume.docx"


def ascii_text(value):
    return value.replace("\u2013", " - ").replace("\u2014", " - ")


def hyperlink(paragraph, label, url):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "174BD6")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def main():
    data = json.loads((ROOT / "src/data/resume.json").read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.03
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "en-US")
    normal.element.get_or_add_rPr().append(language)
    for name, size in [("Title", 25), ("Subtitle", 12), ("Heading 1", 12), ("Heading 2", 11)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].font.bold = True
    # Word's default Title style can carry a blue paragraph border.
    for element in doc.styles.element.iter():
        for child in list(element):
            if child.tag == qn("w:pBdr"):
                element.remove(child)
    bullet = doc.styles["List Bullet"].paragraph_format
    bullet.space_after = Pt(3)
    bullet.line_spacing = 1.03

    doc.add_paragraph(data["name"], "Title")
    doc.add_paragraph(data["headline"], "Subtitle")
    contact = doc.add_paragraph()
    hyperlink(contact, data["email"], "mailto:" + data["email"])
    contact.add_run("  |  ")
    hyperlink(contact, "Portfolio", data["website"])

    doc.add_paragraph("Work Experience", "Heading 1")
    experience = data["experience"]
    p = doc.add_paragraph()
    p.add_run(experience["title"] + " at " + experience["company"]).bold = True
    p.add_run("  |  " + ascii_text(experience["dates"]))
    doc.add_paragraph(experience["location"])
    for item in experience["responsibilities"]:
        doc.add_paragraph(item, "List Bullet")

    doc.add_paragraph("Education", "Heading 1")
    education = data["education"]
    p = doc.add_paragraph()
    p.add_run(education["school"]).bold = True
    doc.add_paragraph(education["degree"] + "  |  Expected " + education["expectedGraduation"])
    doc.add_paragraph(ascii_text(education["dates"]))

    doc.add_paragraph("Technical Skills", "Heading 1")
    for group in data["skills"]:
        p = doc.add_paragraph()
        p.add_run(group["category"] + ": ").bold = True
        p.add_run(", ".join(group["items"]))

    doc.add_paragraph("Project", "Heading 1")
    project = data["project"]
    p = doc.add_paragraph()
    p.add_run(project["title"]).bold = True
    p.add_run("  |  " + ascii_text(project["dates"]))
    doc.add_paragraph(project["summary"])
    doc.add_paragraph("Languages", "Heading 1")
    doc.add_paragraph(", ".join(data["languages"]))

    props = doc.core_properties
    props.title = data["name"] + " Resume"
    props.subject = data["headline"]
    props.author = props.last_modified_by = props.comments = props.keywords = ""
    props.category = props.identifier = props.content_status = props.version = ""
    props.language = "en-US"
    props.created = props.modified = datetime.fromisoformat(data["reviewedAt"]).replace(tzinfo=timezone.utc)
    props.revision = 1
    for element in doc.element.iter():
        for attribute in list(element.attrib):
            if attribute.rsplit("}", 1)[-1].startswith("rsid"):
                del element.attrib[attribute]
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print("Created editable public resume:", OUTPUT)


if __name__ == "__main__":
    main()
